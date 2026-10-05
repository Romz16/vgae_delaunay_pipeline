from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "outputs_revision" / "final_report" / "RELATORIO_INTEGRADO_EVOLUCAO_HIPOTESES_20260919.md"
DOCX_PATH = ROOT / "outputs_revision" / "final_report" / "Relatorio_Integrado_Evolucao_Hipoteses_20260919.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_border(paragraph_or_style) -> None:
    element = paragraph_or_style._element
    p_pr = element.find(qn("w:pPr"))
    if p_pr is None:
        return
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_inline(paragraph, text: str, base_size=10.5, color="222222") -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=base_size - 0.5, color="333333")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=color)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    remove_paragraph_border(title)

    for style_name, size, before, after in (
        ("Heading 1", 17, 15, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=8, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    if rows[0] and rows[0][0] == "Dataset":
        doc.add_page_break()
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    for i, source_row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = source_row[j] if j < len(source_row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i == 0:
                set_cell_shading(cell, "1F4E78")
                run = p.add_run(value)
                set_run_font(run, size=8.5, bold=True, color="FFFFFF")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if i % 2 == 0:
                    set_cell_shading(cell, "F2F6FA")
                add_inline(p, value, base_size=8.4)
                if re.fullmatch(r"[+\-−]?\d+[\d.,% /×^-]*", value.strip()):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_docx() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    style_document(doc)

    footer = doc.sections[0].footer.paragraphs[0]
    add_page_number(footer)

    # Cover page
    title_text = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(70)
    p.add_run(title_text)
    remove_paragraph_border(p)
    for run in p.runs:
        set_run_font(run, name="Aptos Display", size=28, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Síntese dos resultados, testes, evolução das hipóteses e justificativas científicas")
    set_run_font(run, name="Aptos Display", size=15, color="333333")

    date_p = doc.add_paragraph()
    add_inline(date_p, "Consolidação em 19 de setembro de 2026", base_size=11, color="555555")
    scope_p = doc.add_paragraph()
    scope_p.paragraph_format.space_before = Pt(18)
    scope_p.paragraph_format.space_after = Pt(0)
    add_inline(
        scope_p,
        "Abrange os 12 datasets reais, comparações SOTA protocol-matched, controles de construtor, estudos sintéticos, experimento de corrupção-recuperação e validação cega sequencial.",
        base_size=11,
        color="333333",
    )
    doc.add_page_break()

    i = 1
    table_rows: list[list[str]] = []
    in_code = False
    code_buffer: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("**Data de consolidação:**") or stripped.startswith("**Escopo:**"):
            i += 1
            continue

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.right_indent = Cm(0.7)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                for idx, code_line in enumerate(code_buffer):
                    if idx:
                        p.add_run().add_break()
                    run = p.add_run(code_line)
                    set_run_font(run, name="Consolas", size=8.7, color="222222")
                in_code = False
                code_buffer = []
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                i += 1
                continue
            table_rows.append(cells)
            next_is_table = i + 1 < len(lines) and lines[i + 1].strip().startswith("|")
            if not next_is_table:
                add_table(doc, table_rows)
                table_rows = []
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if heading_match:
            hashes, heading = heading_match.groups()
            level = min(len(hashes) - 1, 3)
            doc.add_heading(heading, level=level)
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_together = True
            add_inline(p, stripped[2:], base_size=11, color="1F4E78")
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet_match or number_match:
            content = (bullet_match or number_match).group(1)
            p = doc.add_paragraph(style="List Bullet" if bullet_match else None)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            if number_match:
                number = stripped.split(".", 1)[0]
                p.paragraph_format.left_indent = Cm(0.65)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                add_inline(p, f"{number}.  {content}")
            else:
                add_inline(p, content)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        add_inline(p, stripped)
        i += 1

    props = doc.core_properties
    props.title = title_text
    props.subject = "Resultados e evolução das hipóteses de rewiring geométrico e híbrido"
    props.author = ""
    props.keywords = "rewiring, grafos, Delaunay, VGAE, GNN, validação cega"

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
